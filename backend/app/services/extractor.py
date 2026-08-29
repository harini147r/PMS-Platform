import io
import re
import os
import json
import httpx
from typing import Dict, Any, List, Optional
import fitz # PyMuPDF
from docx import Document

SKILLS_DICTIONARY = [
    "python", "java", "c++", "c", "c#", "javascript", "typescript", "golang", "rust", "ruby", "php", "swift", "kotlin",
    "react", "react.js", "next.js", "vue.js", "angular", "node.js", "express", "fastapi", "django", "flask", "spring boot",
    "sql", "postgresql", "mysql", "mongodb", "redis", "cassandra", "sqlite", "oracle",
    "html", "css", "tailwind css", "bootstrap", "sass", "redux", "graphql", "rest api", "restful apis",
    "docker", "kubernetes", "aws", "azure", "gcp", "ci/cd", "git", "github", "linux", "terraform", "jenkins",
    "machine learning", "deep learning", "nlp", "computer vision", "tensorflow", "pytorch", "scikit-learn", "pandas", "numpy",
    "data structures", "algorithms", "oops", "system design", "microservices", "unit testing", "agile", "scrum",
    "cloud computing", "cybersecurity", "power bi", "tableau", "excel", "data analysis", "devops", "embedded systems", "iot"
]

def extract_text_from_pdf_bytes(pdf_bytes: bytes) -> str:
    """Extract full text from PDF bytes using PyMuPDF."""
    text = ""
    try:
        with fitz.open(stream=pdf_bytes, filetype="pdf") as doc:
            for page in doc:
                text += page.get_text() + "\n"
    except Exception as e:
        print(f"Error parsing PDF: {e}")
    return text.strip()

def extract_text_from_docx_bytes(docx_bytes: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    text = ""
    try:
        doc = Document(io.BytesIO(docx_bytes))
        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"Error parsing DOCX: {e}")
    return text.strip()

async def extract_text_from_url(url: str) -> str:
    """Fetch text or document directly from a public URL or Google Drive link."""
    if not url:
        return ""
    try:
        # Convert standard Google Drive view/share link to direct export or download link
        if "drive.google.com" in url:
            file_id_match = re.search(r"/d/([a-zA-Z0-9_-]+)", url) or re.search(r"id=([a-zA-Z0-9_-]+)", url)
            if file_id_match:
                file_id = file_id_match.group(1)
                url = f"https://drive.google.com/uc?export=download&id={file_id}"
        
        async with httpx.AsyncClient(follow_redirects=True, timeout=15.0) as client:
            resp = await client.get(url)
            if resp.status_code == 200:
                content_type = resp.headers.get("content-type", "").lower()
                content_bytes = resp.content
                if "pdf" in content_type or url.lower().endswith(".pdf"):
                    return extract_text_from_pdf_bytes(content_bytes)
                elif "docx" in content_type or "word" in content_type or url.lower().endswith(".docx"):
                    return extract_text_from_docx_bytes(content_bytes)
                else:
                    text = resp.text
                    clean_text = re.sub(r"<[^>]+>", " ", text)
                    return " ".join(clean_text.split())
    except Exception as e:
        print(f"Error extracting text from URL {url}: {e}")
    return ""

def extract_structured_jd(raw_text: str) -> Dict[str, Any]:
    text_lower = raw_text.lower()
    
    # 1. Extract Role
    role = "Software Development Engineer"
    role_patterns = [
        r"(?:job title|role|position|designation|profile)\s*[:\-\n]\s*([^\n\r]+)",
        r"(?:hiring for|looking for(?: an?)?)\s+([^\n\r,\.]+)",
        r"^(?:job description for|jd -)\s*([^\n\r]+)"
    ]
    for pattern in role_patterns:
        match = re.search(pattern, raw_text, re.IGNORECASE | re.MULTILINE)
        if match:
            extracted = match.group(1).strip()
            if len(extracted) < 60:
                role = extracted.title()
                break
    
    # 2. Extract Skills
    found_skills = []
    for skill in SKILLS_DICTIONARY:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, text_lower):
            found_skills.append(skill.title() if len(skill) > 3 else skill.upper())
    
    unique_skills = []
    for s in found_skills:
        if s not in unique_skills:
            unique_skills.append(s)
    if not unique_skills:
        unique_skills = ["Data Structures", "Algorithms", "Java", "Python", "SQL", "Problem Solving"]

    # 3. Extract Experience
    experience = "0 - 2 Years (Freshers eligible)"
    exp_match = re.search(r"(\d+\s*[\-\+to]+\s*\d*\s*years?(?:\s*of\s*experience)?)", raw_text, re.IGNORECASE)
    if exp_match:
        experience = exp_match.group(1).strip()

    # 4. Extract Qualifications
    qualifications = "B.Tech / B.E / M.Tech / MCA in CSE, IT, ECE or related branches with minimum 65% aggregate."
    if "b.tech" in text_lower or "b.e" in text_lower or "mca" in text_lower or "degree" in text_lower:
        qual_match = re.search(r"(?:qualifications?|eligibility|education)\s*[:\-\n]([\s\S]{20,250}?)(?:\n\n|\n[A-Z]|$)", raw_text, re.IGNORECASE)
        if qual_match:
            qualifications = qual_match.group(1).strip()

    # 5. Responsibilities
    responsibilities = "Collaborate with cross-functional engineering teams, design and build scalable modules, participate in code reviews, and deliver high-quality software."
    resp_match = re.search(r"(?:responsibilities|key tasks|what you(?:'ll| will) do)\s*[:\-\n]([\s\S]{30,400}?)(?:\n\n|\n[A-Z]|$)", raw_text, re.IGNORECASE)
    if resp_match:
        responsibilities = resp_match.group(1).strip()

    return {
        "extracted_role": role,
        "required_skills": unique_skills,
        "qualifications": qualifications,
        "experience": experience,
        "responsibilities": responsibilities
    }
